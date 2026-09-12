"""Tests de la extracción estructural del .docx (Etapa 1 del Word→HTML)."""
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

from app.tools.html_converter.structure import (
    apply_shade,
    apply_tint,
    docx_to_html,
    extract_docx_structure,
)

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _save(doc: Document, tmp_path: Path, name: str) -> Path:
    path = tmp_path / name
    doc.save(str(path))
    return path


# ---------------------------------------------------------------------------
# Tint / shade (matemática HSL, aislada)
# ---------------------------------------------------------------------------
def test_apply_tint_lightens_toward_white():
    assert apply_tint("#000000", 0.5) == "#808080"
    assert apply_tint("#000000", 0.0) == "#FFFFFF"
    assert apply_tint("#000000", 1.0) == "#000000"


def test_apply_shade_darkens_toward_black():
    assert apply_shade("#FFFFFF", 0.5) == "#808080"
    assert apply_shade("#FFFFFF", 0.0) == "#000000"
    assert apply_shade("#FFFFFF", 1.0) == "#FFFFFF"


# ---------------------------------------------------------------------------
# colspan (gridSpan)
# ---------------------------------------------------------------------------
def test_colspan_from_gridspan(tmp_path: Path):
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    c0, c1 = table.rows[0].cells
    c0.text = "Combinada"
    tc_pr = c0._tc.get_or_add_tcPr()
    grid_span = OxmlElement("w:gridSpan")
    grid_span.set(qn("w:val"), "2")
    tc_pr.append(grid_span)
    c1._tc.getparent().remove(c1._tc)
    path = _save(doc, tmp_path, "colspan.docx")

    html = docx_to_html(path)

    assert '<td colspan="2"><p>Combinada</p></td>' in html


# ---------------------------------------------------------------------------
# rowspan (vMerge: restart + N continue)
# ---------------------------------------------------------------------------
def test_rowspan_from_vmerge(tmp_path: Path):
    doc = Document()
    table = doc.add_table(rows=3, cols=1)
    cells = [table.rows[i].cells[0] for i in range(3)]
    cells[0].text = "Vertical"
    tc_pr = cells[0]._tc.get_or_add_tcPr()
    restart = OxmlElement("w:vMerge")
    restart.set(qn("w:val"), "restart")
    tc_pr.append(restart)
    for cell in cells[1:]:
        cell._tc.get_or_add_tcPr().append(OxmlElement("w:vMerge"))
    path = _save(doc, tmp_path, "rowspan.docx")

    html = docx_to_html(path)

    assert '<td rowspan="3"><p>Vertical</p></td>' in html
    assert html.count("<td") == 1


# ---------------------------------------------------------------------------
# Columnas simuladas con tabs
# ---------------------------------------------------------------------------
def test_tab_columns_split_into_cells(tmp_path: Path):
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    p = table.rows[0].cells[0].paragraphs[0]
    p.add_run("Nombre:")
    p.add_run().add_tab()
    p.add_run("Juan Pérez")
    path = _save(doc, tmp_path, "tabs.docx")

    html = docx_to_html(path)

    assert "<td><p>Nombre:</p></td><td><p>Juan Pérez</p></td>" in html
    assert "\t" not in html


# ---------------------------------------------------------------------------
# Colores de tema (resolución end-to-end, leída del theme1.xml del propio doc)
# ---------------------------------------------------------------------------
def _read_theme_color(docx_path: Path, scheme_name: str) -> str:
    with zipfile.ZipFile(docx_path) as zf:
        root = etree.fromstring(zf.read("word/theme/theme1.xml"))
    scheme = root.find(f".//{{{A_NS}}}clrScheme")
    element = scheme.find(f"{{{A_NS}}}{scheme_name}")
    srgb = element.find(f"{{{A_NS}}}srgbClr")
    return "#" + srgb.get("val").upper()


def test_theme_color_resolution(tmp_path: Path):
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.text = "color"
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:themeColor"), "accent1")
    tc_pr.append(shd)
    path = _save(doc, tmp_path, "theme.docx")

    expected = _read_theme_color(path, "accent1")
    structure = extract_docx_structure(path)
    assert structure.blocks[0].rows[0].cells[0].background == expected

    html = docx_to_html(path)
    assert f'style="background-color:{expected}"' in html


# ---------------------------------------------------------------------------
# Color directo (w:fill)
# ---------------------------------------------------------------------------
def test_direct_fill_color(tmp_path: Path):
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.text = "x"
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "FF0000")
    tc_pr.append(shd)
    path = _save(doc, tmp_path, "fill.docx")

    html = docx_to_html(path)

    assert 'style="background-color:#FF0000"' in html


# ---------------------------------------------------------------------------
# Anchos de columna (tblGrid, twips -> porcentaje)
# ---------------------------------------------------------------------------
def test_column_widths_percentages(tmp_path: Path):
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    grid = table._tbl.find(qn("w:tblGrid"))
    cols = grid.findall(qn("w:gridCol"))
    cols[0].set(qn("w:w"), "1440")
    cols[1].set(qn("w:w"), "2880")
    path = _save(doc, tmp_path, "widths.docx")

    html = docx_to_html(path)

    assert '<colgroup><col style="width:33.33%"><col style="width:66.67%"></colgroup>' in html


# ---------------------------------------------------------------------------
# Tipografía por defecto (del estilo Normal, no hardcodeado)
# ---------------------------------------------------------------------------
def test_font_defaults_extracted(tmp_path: Path):
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(14)
    doc.add_paragraph("hola")
    path = _save(doc, tmp_path, "font.docx")

    structure = extract_docx_structure(path)

    assert structure.default_font_family == "Arial"
    assert structure.default_font_size_pt == 14.0


# ---------------------------------------------------------------------------
# Cero &nbsp;
# ---------------------------------------------------------------------------
def test_no_nbsp_in_output(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("texto normal")
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = ""
    path = _save(doc, tmp_path, "nbsp.docx")

    html = docx_to_html(path)

    assert "\xa0" not in html
    assert "&nbsp;" not in html


# ---------------------------------------------------------------------------
# Subrayado (issue #58): el w:val de <w:u> decide, no su presencia
# ---------------------------------------------------------------------------
def test_underline_detection_uses_w_val(tmp_path: Path):
    """w:val="none" y la ausencia de <w:u> no deben dar <u>; solo single sí."""
    doc = Document()
    doc.add_paragraph("sin etiqueta w:u")  # nunca se toca underline → sin <w:u>

    p_none = doc.add_paragraph()
    r_none = p_none.add_run("val none")
    r_none.underline = False  # Word escribe <w:u w:val="none"/>

    p_single = doc.add_paragraph()
    r_single = p_single.add_run("val single")
    r_single.underline = True  # <w:u w:val="single"/>

    path = _save(doc, tmp_path, "underline.docx")

    html = docx_to_html(path)

    assert "<p>sin etiqueta w:u</p>" in html
    assert "<p>val none</p>" in html
    assert "<p><u>val single</u></p>" in html


# ---------------------------------------------------------------------------
# Negrita e itálica (issue #62): el w:val de <w:b>/<w:i> decide, no su presencia
# ---------------------------------------------------------------------------
def test_bold_detection_uses_w_val(tmp_path: Path):
    """w:val="0" y la ausencia de <w:b> no deben dar <strong>; bare sí."""
    doc = Document()
    doc.add_paragraph("sin etiqueta w:b")  # nunca se toca bold → sin <w:b>

    p_off = doc.add_paragraph()
    p_off.add_run("val 0").bold = False  # <w:b w:val="0"/>

    p_on = doc.add_paragraph()
    p_on.add_run("bare").bold = True  # <w:b/> (sin w:val)

    path = _save(doc, tmp_path, "bold.docx")

    html = docx_to_html(path)

    assert "<p>sin etiqueta w:b</p>" in html
    assert "<p>val 0</p>" in html
    assert "<p><strong>bare</strong></p>" in html


def test_italic_detection_uses_w_val(tmp_path: Path):
    """w:val="0" y la ausencia de <w:i> no deben dar <em>; bare sí."""
    doc = Document()
    doc.add_paragraph("sin etiqueta w:i")  # nunca se toca italic → sin <w:i>

    p_off = doc.add_paragraph()
    p_off.add_run("val 0").italic = False  # <w:i w:val="0"/>

    p_on = doc.add_paragraph()
    p_on.add_run("bare").italic = True  # <w:i/> (sin w:val)

    path = _save(doc, tmp_path, "italic.docx")

    html = docx_to_html(path)

    assert "<p>sin etiqueta w:i</p>" in html
    assert "<p>val 0</p>" in html
    assert "<p><em>bare</em></p>" in html
