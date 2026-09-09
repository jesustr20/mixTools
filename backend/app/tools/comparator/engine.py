"""
Herramienta 3: Comparador de documentos (PDF/Word/Excel en cualquier combinación).

Draftable hace diff VISUAL (superpone páginas como imágenes) + diff de texto.
Aquí implementamos primero el diff de TEXTO a nivel de palabra (igual al modo
"inline" de Draftable: rojo tachado = se quitó, verde subrayado = se agregó).
El diff visual (superposición de páginas PDF) queda apuntado como Fase 2 al
final de este archivo, porque requiere alinear páginas imagen a imagen.
"""
import difflib
import html
from pathlib import Path

import openpyxl
import pymupdf as fitz  # PyMuPDF (nombre nuevo, "fitz" queda deprecado)
from docx import Document


# ---------------------------------------------------------------------------
# Extracción de texto por tipo de archivo
# ---------------------------------------------------------------------------
def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(path)
    elif ext == ".docx":
        return _extract_docx(path)
    elif ext in (".xlsx", ".xlsm"):
        return _extract_xlsx(path)
    else:
        raise ValueError(f"Tipo de archivo no soportado para comparar: {ext}")


def _extract_pdf(path: Path) -> str:
    doc = fitz.open(path)
    text = "\n".join(page.get_text() for page in doc)
    doc.close()
    return text


def _extract_docx(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_xlsx(path: Path) -> str:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"# Hoja: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Diff de texto a nivel de palabra -> HTML con resaltado (estilo Draftable)
# ---------------------------------------------------------------------------
def diff_to_html(text_a: str, text_b: str) -> str:
    words_a = text_a.split()
    words_b = text_b.split()

    matcher = difflib.SequenceMatcher(a=words_a, b=words_b, autojunk=False)
    out = []
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            out.append(html.escape(" ".join(words_a[i1:i2])))
        elif tag == "delete":
            out.append(f'<del class="diff-del">{html.escape(" ".join(words_a[i1:i2]))}</del>')
        elif tag == "insert":
            out.append(f'<ins class="diff-ins">{html.escape(" ".join(words_b[j1:j2]))}</ins>')
        elif tag == "replace":
            out.append(f'<del class="diff-del">{html.escape(" ".join(words_a[i1:i2]))}</del>')
            out.append(f'<ins class="diff-ins">{html.escape(" ".join(words_b[j1:j2]))}</ins>')
    return " ".join(out)


def diff_stats(text_a: str, text_b: str) -> dict:
    words_a = text_a.split()
    words_b = text_b.split()
    matcher = difflib.SequenceMatcher(a=words_a, b=words_b, autojunk=False)
    added = deleted = 0
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag in ("delete", "replace"):
            deleted += i2 - i1
        if tag in ("insert", "replace"):
            added += j2 - j1
    return {"palabras_agregadas": added, "palabras_eliminadas": deleted,
            "similitud": round(matcher.ratio() * 100, 1)}


def compare_files(path_a: Path, path_b: Path) -> dict:
    text_a = extract_text(path_a)
    text_b = extract_text(path_b)
    return {
        "diff_html": diff_to_html(text_a, text_b),
        "stats": diff_stats(text_a, text_b),
    }

# ---------------------------------------------------------------------------
# FASE 2 (pendiente, no implementada todavía):
# Diff VISUAL real como Draftable -> renderizar cada página de ambos PDFs
# a imagen (fitz), alinear página N vs página N (o usar un algoritmo de
# alineación si difiere el número de páginas), y superponer con diferencia
# de píxeles o contornos de texto movido/añadido/eliminado resaltados en
# capas de color. Requiere más trabajo (alineación robusta) — lo dejamos
# como siguiente iteración una vez que el diff de texto esté validado.
# ---------------------------------------------------------------------------
