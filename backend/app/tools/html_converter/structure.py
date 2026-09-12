"""
Extracción estructural del XML de un .docx (Etapa 1 del Word→HTML).

Leemos el XML directamente vía python-docx (no un conversor genérico) para
recuperar datos que un HTML simplificado descarta: colores de celda
(incluyendo colores de tema con tint/shade), anchos de columna reales
(twips→porcentaje) y colspan/rowspan reales (gridSpan/vMerge). Sin IA
(eso es la Etapa 2).

La salida es un árbol de dataclasses (`DocumentStructure`) y una función de
render que lo convierte en HTML. El esqueleto de salida fijo (cabecera, pie,
Liquid) se agrega en etapas posteriores; acá solo se convierte el contenido.
"""
import colorsys
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

# w:themeColor -> nombre del elemento dentro de <a:clrScheme> en theme1.xml
_THEME_COLOR_TO_SCHEME = {
    "text1": "dk1",
    "background1": "lt1",
    "text2": "dk2",
    "background2": "lt2",
    "accent1": "accent1",
    "accent2": "accent2",
    "accent3": "accent3",
    "accent4": "accent4",
    "accent5": "accent5",
    "accent6": "accent6",
    "hyperlink": "hlink",
    "followedHyperlink": "folHlink",
}


# ---------------------------------------------------------------------------
# Modelo (árbol de dataclasses)
# ---------------------------------------------------------------------------
@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    underline: bool = False


@dataclass
class Paragraph:
    runs: list[Run] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "".join(r.text for r in self.runs)


@dataclass
class Cell:
    paragraphs: list[Paragraph] = field(default_factory=list)
    colspan: int = 1
    rowspan: int = 1
    background: str | None = None

    @property
    def text(self) -> str:
        return "\n".join(p.text for p in self.paragraphs)


@dataclass
class Row:
    cells: list[Cell] = field(default_factory=list)


@dataclass
class Table:
    rows: list[Row] = field(default_factory=list)
    column_widths_pct: list[float] = field(default_factory=list)


@dataclass
class DocumentStructure:
    blocks: list[Paragraph | Table] = field(default_factory=list)
    default_font_family: str | None = None
    default_font_size_pt: float | None = None


# ---------------------------------------------------------------------------
# Tint / shade (matemática HSL, aislada para poder testearla en unidad)
# ---------------------------------------------------------------------------
def apply_tint(hex_color: str, tint: float) -> str:
    """Aclara el color hacia blanco. `tint` en [0,1]; 1.0 = sin cambio."""
    return _adjust_luminance(hex_color, lambda lum: lum * tint + (1.0 - tint))


def apply_shade(hex_color: str, shade: float) -> str:
    """Oscurece el color hacia negro. `shade` en [0,1]; 1.0 = sin cambio."""
    return _adjust_luminance(hex_color, lambda lum: lum * shade)


def _adjust_luminance(hex_color: str, transform) -> str:
    r, g, b = _hex_to_rgb01(hex_color)
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = max(0.0, min(1.0, transform(l)))
    r2, g2, b2 = colorsys.hls_to_rgb(h, l, s)
    return _rgb01_to_hex(r2, g2, b2)


def _hex_to_rgb01(hex_color: str) -> tuple[float, float, float]:
    v = hex_color.strip().lstrip("#")
    return (int(v[0:2], 16) / 255.0, int(v[2:4], 16) / 255.0, int(v[4:6], 16) / 255.0)


def _rgb01_to_hex(r: float, g: float, b: float) -> str:
    def _c(v: float) -> int:
        return max(0, min(255, round(v * 255)))

    return f"#{_c(r):02X}{_c(g):02X}{_c(b):02X}"


def _normalize_hex(value: str) -> str:
    v = value.strip().lstrip("#")
    return "#" + v[:6].upper()


# ---------------------------------------------------------------------------
# Tema (theme1.xml) y resolución de colores de tema
# ---------------------------------------------------------------------------
def _load_theme_colors(docx_path: Path) -> dict[str, str]:
    """Devuelve {nombre_clrScheme: hex} (ej. {"accent1": "4F81BD"})."""
    try:
        with zipfile.ZipFile(docx_path) as zf:
            xml = zf.read("word/theme/theme1.xml")
    except (KeyError, zipfile.BadZipFile):
        return {}

    root = etree.fromstring(xml)
    scheme = root.find(f".//{{{A_NS}}}clrScheme")
    if scheme is None:
        return {}

    colors: dict[str, str] = {}
    for child in scheme:
        name = etree.QName(child).localname
        srgb = child.find(f"{{{A_NS}}}srgbClr")
        if srgb is not None and srgb.get("val"):
            colors[name] = srgb.get("val")
        else:
            sys_clr = child.find(f"{{{A_NS}}}sysClr")
            if sys_clr is not None and sys_clr.get("lastClr"):
                colors[name] = sys_clr.get("lastClr")
    return colors


def _resolve_shading(shd, theme_colors: dict[str, str]) -> str | None:
    theme_color = shd.get(qn("w:themeColor"))
    if theme_color:
        scheme_name = _THEME_COLOR_TO_SCHEME.get(theme_color)
        base = theme_colors.get(scheme_name) if scheme_name else None
        if base is None:
            return None
        tint = shd.get(qn("w:themeTint"))
        shade = shd.get(qn("w:themeShade"))
        if tint is not None:
            return apply_tint(_normalize_hex(base), int(tint, 16) / 255.0)
        if shade is not None:
            return apply_shade(_normalize_hex(base), int(shade, 16) / 255.0)
        return _normalize_hex(base)

    fill = shd.get(qn("w:fill"))
    if fill and fill.lower() != "auto":
        return _normalize_hex(fill)
    return None


# ---------------------------------------------------------------------------
# Extracción
# ---------------------------------------------------------------------------
def extract_docx_structure(docx_path: Path) -> DocumentStructure:
    document = Document(str(docx_path))
    theme_colors = _load_theme_colors(docx_path)

    structure = DocumentStructure()
    normal = document.styles["Normal"]
    if normal.font.name is not None:
        structure.default_font_family = normal.font.name
    if normal.font.size is not None:
        structure.default_font_size_pt = normal.font.size.pt

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            structure.blocks.append(_parse_paragraph(child))
        elif child.tag == qn("w:tbl"):
            structure.blocks.append(_parse_table(child, theme_colors))

    return structure


def _parse_paragraph(p_elem) -> Paragraph:
    runs = [_parse_run(r) for r in p_elem.iter(qn("w:r"))]
    return Paragraph(runs=runs)


def _parse_run(r_elem) -> Run:
    parts: list[str] = []
    for node in r_elem:
        if node.tag == qn("w:t"):
            parts.append(node.text or "")
        elif node.tag == qn("w:tab"):
            parts.append("\t")
        elif node.tag in (qn("w:br"), qn("w:cr")):
            parts.append("\n")

    r_pr = r_elem.find(qn("w:rPr"))
    bold = italic = underline = False
    if r_pr is not None:
        bold = _is_flag_active(r_pr, "w:b")
        italic = _is_flag_active(r_pr, "w:i")
        underline = _is_flag_active(r_pr, "w:u")

    return Run(text="".join(parts), bold=bold, italic=italic, underline=underline)


def _is_flag_active(r_pr, tag_name: str) -> bool:
    """True si la propiedad de run `tag_name` está activa, mirando su w:val.

    Word escribe <w:b>/<w:i>/<w:u> con un w:val explícito incluso para
    desactivarlos: negrita e itálica usan "0"/"false" (inactivo) y "1"/"true" o
    elemento sin w:val (activo); subrayado usa "none" (inactivo) y cualquier
    estilo, p. ej. "single" (activo). La ausencia del elemento o un w:val que
    desactiva explícitamente dan False (issues #58 y #62).
    """
    el = r_pr.find(qn(tag_name))
    if el is None:
        return False
    return el.get(qn("w:val")) not in ("0", "false", "none")


def _parse_table(tbl_elem, theme_colors: dict[str, str]) -> Table:
    column_widths_pct = _column_widths_pct(tbl_elem)
    grid = [_row_cells_with_grid(tr) for tr in tbl_elem.findall(qn("w:tr"))]

    rows: list[Row] = []
    for r_idx, row_cells in enumerate(grid):
        cells: list[Cell] = []
        for start, _end, tc in row_cells:
            if _vmerge_state(tc) == "continue":
                continue
            colspan = _end - start
            rowspan = 1
            if _vmerge_state(tc) == "restart":
                rowspan = _count_rowspan(grid, r_idx, start)
            cells.append(
                Cell(
                    paragraphs=_cell_paragraphs(tc),
                    colspan=colspan,
                    rowspan=rowspan,
                    background=_cell_background(tc, theme_colors),
                )
            )
        rows.append(Row(cells=_split_tab_cells(cells)))

    return Table(rows=rows, column_widths_pct=column_widths_pct)


def _column_widths_pct(tbl_elem) -> list[float]:
    grid = tbl_elem.find(qn("w:tblGrid"))
    if grid is None:
        return []
    widths = []
    for gc in grid.findall(qn("w:gridCol")):
        try:
            widths.append(int(gc.get(qn("w:w")) or 0))
        except ValueError:
            widths.append(0)
    total = sum(widths)
    if total <= 0:
        return []
    return [round(w / total * 100, 2) for w in widths]


def _row_cells_with_grid(tr_elem) -> list[tuple[int, int, etree._Element]]:
    result: list[tuple[int, int, etree._Element]] = []
    col = 0
    for tc in tr_elem.findall(qn("w:tc")):
        span = _grid_span(tc)
        result.append((col, col + span, tc))
        col += span
    return result


def _grid_span(tc) -> int:
    tc_pr = tc.find(qn("w:tcPr"))
    if tc_pr is None:
        return 1
    gs = tc_pr.find(qn("w:gridSpan"))
    if gs is None:
        return 1
    try:
        return int(gs.get(qn("w:val")))
    except (TypeError, ValueError):
        return 1


def _vmerge_state(tc) -> str | None:
    tc_pr = tc.find(qn("w:tcPr"))
    if tc_pr is None:
        return None
    vm = tc_pr.find(qn("w:vMerge"))
    if vm is None:
        return None
    return "restart" if vm.get(qn("w:val")) == "restart" else "continue"


def _count_rowspan(grid, r_idx: int, start_col: int) -> int:
    rowspan = 1
    for rr in range(r_idx + 1, len(grid)):
        continues = False
        for s2, e2, tc2 in grid[rr]:
            if s2 <= start_col < e2:
                if _vmerge_state(tc2) == "continue":
                    rowspan += 1
                    continues = True
                break
        if not continues:
            break
    return rowspan


def _cell_paragraphs(tc) -> list[Paragraph]:
    return [_parse_paragraph(p) for p in tc.findall(qn("w:p"))]


def _cell_background(tc, theme_colors: dict[str, str]) -> str | None:
    tc_pr = tc.find(qn("w:tcPr"))
    if tc_pr is None:
        return None
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        return None
    return _resolve_shading(shd, theme_colors)


def _split_tab_cells(cells: list[Cell]) -> list[Cell]:
    """Reconstruye "columnas simuladas con tabs" como celdas reales.

    Heurística de la Etapa 1: una celda cuyo texto contiene un tab (\t) se
    parte en N celdas (una por segmento). Útil para tablas que en el .docx
    usan una sola celda con tabs para fingir dos columnas.
    """
    result: list[Cell] = []
    for cell in cells:
        parts = cell.text.split("\t")
        if len(parts) == 1:
            result.append(cell)
            continue
        for part in parts:
            result.append(
                Cell(
                    paragraphs=[Paragraph(runs=[Run(text=part)])],
                    colspan=1,
                    rowspan=1,
                    background=cell.background,
                )
            )
    return result


# ---------------------------------------------------------------------------
# Render
# ---------------------------------------------------------------------------
def render_structure_html(structure: DocumentStructure) -> str:
    return "".join(
        _render_table(block) if isinstance(block, Table) else _render_paragraph(block)
        for block in structure.blocks
    )


def _render_paragraph(p: Paragraph) -> str:
    if not p.runs:
        return "<p></p>"
    inner = "".join(_render_run(r) for r in p.runs)
    return f"<p>{inner}</p>"


def _render_run(r: Run) -> str:
    text = _escape_html(r.text).replace("\n", "<br>")
    if r.bold:
        text = f"<strong>{text}</strong>"
    if r.italic:
        text = f"<em>{text}</em>"
    if r.underline:
        text = f"<u>{text}</u>"
    return text


def _render_table(table: Table) -> str:
    parts = ["<table>"]
    if table.column_widths_pct:
        cols = "".join(
            f'<col style="width:{w:.2f}%">' for w in table.column_widths_pct
        )
        parts.append(f"<colgroup>{cols}</colgroup>")
    for row in table.rows:
        parts.append("<tr>")
        for cell in row.cells:
            parts.append(_render_cell(cell))
        parts.append("</tr>")
    parts.append("</table>")
    return "".join(parts)


def _render_cell(cell: Cell) -> str:
    attrs: list[str] = []
    if cell.colspan > 1:
        attrs.append(f'colspan="{cell.colspan}"')
    if cell.rowspan > 1:
        attrs.append(f'rowspan="{cell.rowspan}"')
    if cell.background:
        attrs.append(f'style="background-color:{cell.background}"')
    attr_str = (" " + " ".join(attrs)) if attrs else ""
    inner = "".join(_render_paragraph(p) for p in cell.paragraphs)
    return f"<td{attr_str}>{inner}</td>"


def _escape_html(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def docx_to_html(docx_path: Path) -> str:
    """Extrae la estructura y la renderiza a HTML en un solo paso."""
    return render_structure_html(extract_docx_structure(docx_path))
